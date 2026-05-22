import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from dotenv import load_dotenv


SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff"}
PICTURE_WIDTH = Inches(3.6)


def natural_sort_key(name: str):
    return [int(tok) if tok.isdigit() else tok.lower()
            for tok in re.split(r"(\d+)", name)]


def collect_images(folder: Path) -> list[Path]:
    images = [p for p in folder.iterdir()
              if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS]
    images.sort(key=lambda p: natural_sort_key(p.name))
    return images


def clear_cell(cell) -> None:
    tc = cell._element
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(OxmlElement("w:p"))


def insert_image_in_cell(cell, image_path: Path) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=PICTURE_WIDTH)


def main() -> None:
    load_dotenv()

    images_folder = os.getenv("INPUT_IMAGES_FOLDER_PATH")
    output_folder = os.getenv("REPORT_INPUT_FORMS_FOLDER")
    template_env = os.getenv("INPUT_FORM_TEMPLATE_PATH")

    if not images_folder:
        sys.exit("INPUT_IMAGES_FOLDER_PATH is not set in .env")
    if not output_folder:
        sys.exit("REPORT_INPUT_FORMS_FOLDER is not set in .env")
    if not template_env:
        sys.exit("INPUT_FORM_TEMPLATE_PATH is not set in .env")

    images_dir = Path(images_folder)
    if not images_dir.is_dir():
        sys.exit(f"Images folder does not exist: {images_dir}")

    images = collect_images(images_dir)
    if not images:
        sys.exit(f"No supported images found in {images_dir}. "
                 f"Supported extensions: {sorted(SUPPORTED_EXTENSIONS)}")

    template_path = Path(template_env)
    if not template_path.is_file():
        sys.exit(f"Input-form template not found: {template_path}")

    output_dir = Path(output_folder)
    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"{images_dir.name}_{timestamp}.docx"

    shutil.copy(template_path, output_path)

    doc = Document(str(output_path))
    if len(doc.tables) < 2:
        sys.exit("Template does not contain a second table — aborting.")
    table = doc.tables[1]

    # Clear every data row (preserve header row 0).
    for row in table.rows[1:]:
        for cell in row.cells:
            clear_cell(cell)

    for i, image_path in enumerate(images):
        row_index = i + 1  # row 0 is the header
        if row_index >= len(table.rows):
            table.add_row()
        picture_cell = table.rows[row_index].cells[2]
        insert_image_in_cell(picture_cell, image_path)

    doc.save(str(output_path))
    print(f"Inserted {len(images)} image(s) into: {output_path}")


if __name__ == "__main__":
    main()
